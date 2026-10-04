#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Generate tests/fixture_vrain_like.pptx: a stand-in with the VRAIN proposal
template's *structure* (not its artwork), so the reference example can be run
end to end without the real, confidential template.

Same shape as the real one: A4; 39 slides; 1 cover, 2 TOC, 3/5/8/11 dividers,
4/6/7/9/10 body slides on a layout that supplies body placeholders idx 12
(title) and 13 (lead message); 12-25 previous-deal result slides with images;
26 Appendix divider; 27-39 appendix slides with images. All names fictional.
Requires python-pptx (only to generate; the committed .pptx needs nothing).
"""
import io
import os
import re
import shutil
import tempfile
import zipfile

from PIL import Image
from pptx import Presentation
from pptx.util import Emu, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "fixture_vrain_like.pptx")
NAVY = (0x11, 0x2C, 0x4B)

LAYOUT_EXTRA = (
    '<p:sp><p:nvSpPr><p:cNvPr id="90" name="TopBar"/><p:cNvSpPr/><p:nvPr userDrawn="1"/></p:nvSpPr>'
    '<p:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="9906000" cy="200742"/></a:xfrm>'
    '<a:prstGeom prst="rect"><a:avLst/></a:prstGeom><a:solidFill><a:srgbClr val="112C4B"/></a:solidFill>'
    '<a:ln><a:noFill/></a:ln></p:spPr></p:sp>'
    '<p:sp><p:nvSpPr><p:cNvPr id="91" name="Rule"/><p:cNvSpPr/><p:nvPr userDrawn="1"/></p:nvSpPr>'
    '<p:spPr><a:xfrm><a:off x="0" y="876676"/><a:ext cx="9906000" cy="11773"/></a:xfrm>'
    '<a:prstGeom prst="rect"><a:avLst/></a:prstGeom><a:solidFill><a:srgbClr val="112C4B"/></a:solidFill>'
    '<a:ln><a:noFill/></a:ln></p:spPr></p:sp>'
    '<p:sp><p:nvSpPr><p:cNvPr id="92" name="Conf"/><p:cNvSpPr txBox="1"/><p:nvPr userDrawn="1"/></p:nvSpPr>'
    '<p:spPr><a:xfrm><a:off x="7234238" y="419239"/><a:ext cx="1600000" cy="246395"/></a:xfrm>'
    '<a:prstGeom prst="rect"><a:avLst/></a:prstGeom><a:noFill/><a:ln w="19050"><a:solidFill>'
    '<a:srgbClr val="FF0000"/></a:solidFill></a:ln></p:spPr><p:txBody><a:bodyPr anchor="ctr"/><a:lstStyle/>'
    '<a:p><a:pPr algn="ctr"/><a:r><a:rPr lang="en-US" sz="900" b="1"><a:solidFill><a:srgbClr val="FF0000"/>'
    '</a:solidFill></a:rPr><a:t>CONFIDENTIAL (fixture)</a:t></a:r></a:p></p:txBody></p:sp>'
)


def ph(idx, x, y, cx, cy, prompt):
    return ('<p:sp><p:nvSpPr><p:cNvPr id="%d" name="PH %d"/><p:cNvSpPr><a:spLocks noGrp="1"/></p:cNvSpPr>'
            '<p:nvPr><p:ph type="body" sz="quarter" idx="%d" hasCustomPrompt="1"/></p:nvPr></p:nvSpPr>'
            '<p:spPr><a:xfrm><a:off x="%d" y="%d"/><a:ext cx="%d" cy="%d"/></a:xfrm></p:spPr>'
            '<p:txBody><a:bodyPr anchor="ctr"/><a:lstStyle><a:lvl1pPr marL="0" indent="0"><a:buNone/>'
            '</a:lvl1pPr></a:lstStyle><a:p><a:r><a:rPr lang="ja-JP"/><a:t>%s</a:t></a:r></a:p></p:txBody></p:sp>'
            % (80 + idx, idx, idx, x, y, cx, cy, prompt))


def tb(slide, x, y, w, h, paras, size=18, bold=False):
    box = slide.shapes.add_textbox(Emu(x), Emu(y), Emu(w), Emu(h))
    tf = box.text_frame
    for k, t in enumerate(paras):
        p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
        if isinstance(t, tuple):           # several runs in one paragraph
            for piece in t:
                r = p.add_run(); r.text = piece; r.font.size = Pt(size); r.font.bold = bold
        else:
            r = p.add_run(); r.text = t; r.font.size = Pt(size); r.font.bold = bold
    return box


def main():
    tmp = tempfile.mkdtemp()
    img = os.path.join(tmp, "img.png")
    Image.new("RGB", (240, 140), NAVY).save(img)

    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(9906000), Emu(6858000)
    blank = prs.slide_layouts[6]
    sections = ["現状のご状況と課題", "ご提案", "ご検討にあたっての比較観点", "検証結果（再掲）"]
    for n in range(1, 40):
        s = prs.slides.add_slide(blank)
        if n == 1:
            tb(s, 307975, 865922, 7195004, 600000, [("サンプル食品株式", "会社　御中")], 16, True)
            tb(s, 1112750, 3034903, 7680487, 900000, ["サンプル容器外観検査（アドオン）", "ご提案書"], 20, True)
            tb(s, 5800000, 5200000, 3600000, 1200000,
               ["2026年1月1日", ("株式会社", "VRAIN Solution"), "ソリューション事業部", "担当 太郎"], 11)
        elif n == 2:
            tb(s, 300000, 300000, 3000000, 500000, ["目次"], 20, True)
            for k, (num, name) in enumerate(zip(["Ⅰ.", "Ⅱ.", "Ⅲ.", "Ⅳ."], sections)):
                tb(s, 1200000, 1500000 + k * 800000, 6000000, 500000, [(num, name)], 14, True)
            tb(s, 1200000, 1500000 + 4 * 800000, 6000000, 500000, [("Ⅴ", ".", " Appendix")], 14, True)
        elif n in (3, 5, 8, 11):
            num = {3: "Ⅰ.", 5: "Ⅱ.", 8: "Ⅲ.", 11: "Ⅳ."}[n]
            tb(s, 1100000, 2700000, 7000000, 600000, [(num, sections[[3, 5, 8, 11].index(n)])], 20, True)
        elif n == 26:
            tb(s, 1100000, 2700000, 7000000, 600000, [("Ⅴ", ".", " Appendix")], 20, True)
        elif 12 <= n <= 25:
            tb(s, 400000, 300000, 8000000, 500000, ["前案件の検証結果 %d" % n], 18, True)
            s.shapes.add_picture(img, Emu(1000000), Emu(2000000))
        elif n >= 27:
            tb(s, 400000, 300000, 8000000, 500000, ["会社紹介 %d" % n], 18, True)
            s.shapes.add_picture(img, Emu(1000000), Emu(2000000))
        else:                               # body slides 4,6,7,9,10
            tb(s, 400000, 2000000, 8000000, 500000, ["前案件の本編 %d" % n], 14)
    pptx_tmp = os.path.join(tmp, "f.pptx")
    prs.save(pptx_tmp)

    # add the VRAIN-like header/placeholders to the blank layout all slides use
    work = os.path.join(tmp, "u")
    zipfile.ZipFile(pptx_tmp).extractall(work)
    lay = os.path.join(work, "ppt", "slideLayouts", "slideLayout7.xml")
    x = io.open(lay, encoding="utf-8").read()
    extra = LAYOUT_EXTRA + ph(12, 125153, 265978, 7000000, 551476, "タイトル") + \
        ph(13, 125153, 1087316, 9655692, 467158, "メッセージ")
    x = re.sub(r'(</p:spTree>)', extra + r'\1', x, count=1)
    io.open(lay, "w", encoding="utf-8").write(x)
    if os.path.exists(OUT):
        os.remove(OUT)
    with zipfile.ZipFile(OUT, "w", zipfile.ZIP_DEFLATED) as z:
        for root, _, files in os.walk(work):
            for f in files:
                full = os.path.join(root, f)
                z.write(full, os.path.relpath(full, work).replace("\\", "/"))
    shutil.rmtree(tmp, ignore_errors=True)
    print("wrote", OUT, os.path.getsize(OUT), "bytes")


if __name__ == "__main__":
    main()
